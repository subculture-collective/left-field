#!/usr/bin/env python3
"""Create a deterministic candidate-total extract from locked California SOV workbooks."""

from __future__ import annotations

import argparse
import csv
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl


EXPECTED_OPENPYXL_VERSION = "3.1.5"
if openpyxl.__version__ != EXPECTED_OPENPYXL_VERSION:
    raise RuntimeError(
        "California extraction requires "
        f"openpyxl {EXPECTED_OPENPYXL_VERSION}; found {openpyxl.__version__}"
    )


DISTRICT_HEADING = re.compile(r"^(\d+)(?:st|nd|rd|th) Congressional District$")
PARTY = re.compile(r"^(?:AI|DEM|GRN|LIB|NPP|PF|REP)(?: \(W/I\))?$")
HEADER = [
    "cycle_year",
    "election_date",
    "district_code",
    "candidate_order",
    "source_candidate_name",
    "party_preference",
    "is_write_in",
    "source_incumbent_marker",
    "votes",
    "supplied_percent_tenths",
]
EXPECTED = {
    2022: {"date": "2022-06-07", "candidates": 272, "votes": 6_896_731},
    2024: {"date": "2024-03-05", "candidates": 245, "votes": 7_283_233},
    2026: {"date": "2026-06-02", "candidates": 297, "votes": 8_827_913},
}


def clean_name(value: object) -> tuple[str, bool]:
    name = " ".join(str(value).replace("\n", " ").split())
    incumbent = name.endswith("*")
    if incumbent:
        name = name[:-1].rstrip()
    if not name:
        raise ValueError("empty candidate name")
    return name, incumbent


def percent_tenths(value: object) -> int:
    if isinstance(value, str):
        text = value.strip()
        if not text.endswith("%") or text == "%":
            raise ValueError(f"invalid supplied percentage {value!r}")
        percentage = Decimal(text[:-1])
    elif isinstance(value, (int, float)):
        percentage = Decimal(str(value)) * 100
    else:
        raise ValueError(f"invalid supplied percentage {value!r}")
    result = int((percentage * 10).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    if not 0 <= result <= 1000:
        raise ValueError(f"percentage outside range {value!r}")
    return result


def candidate_row(year: int, district: int, order: int, name_value: object, party_value: object, votes_value: object, percent_value: object) -> list[object]:
    name, incumbent = clean_name(name_value)
    party = str(party_value).strip()
    if not PARTY.fullmatch(party):
        raise ValueError(f"invalid party preference {party!r}")
    if not isinstance(votes_value, int) or votes_value < 0:
        raise ValueError(f"invalid vote total {votes_value!r}")
    is_write_in = party.endswith(" (W/I)")
    return [year, EXPECTED[year]["date"], f"{district:02d}", order, name, party.removesuffix(" (W/I)"), int(is_write_in), int(incumbent), votes_value, percent_tenths(percent_value)]


def formatted_sheet_rows(path: Path, year: int) -> list[list[object]]:
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True)["Representative in Congress"]
    rows = list(sheet.iter_rows(values_only=True))
    starts: list[tuple[int, int]] = []
    for index, row in enumerate(rows):
        match = DISTRICT_HEADING.fullmatch(str(row[0]).strip()) if row and row[0] is not None else None
        if match:
            starts.append((index, int(match.group(1))))
    if [district for _, district in starts] != list(range(1, 53)):
        raise ValueError(f"{year}: district heading closure invalid")

    output: list[list[object]] = []
    for position, (start, district) in enumerate(starts):
        end = starts[position + 1][0] if position + 1 < len(starts) else len(rows)
        block = rows[start + 1 : end]
        order = 0
        for index in range(len(block) - 1):
            names, parties = block[index], block[index + 1]
            columns = [column for column in range(1, len(parties)) if parties[column] is not None]
            if not columns or not all(PARTY.fullmatch(str(parties[column]).strip()) for column in columns):
                continue
            if not all(names[column] is not None for column in columns):
                raise ValueError(f"{year} district {district}: party column without candidate")
            total_row = percent_row = None
            for later in block[index + 2 :]:
                label = str(later[0]).strip() if later[0] is not None else ""
                if label == "District Totals":
                    total_row = later
                elif label == "Percent" and total_row is not None:
                    percent_row = later
                    break
            if total_row is None or percent_row is None:
                raise ValueError(f"{year} district {district}: totals missing")
            for column in columns:
                order += 1
                output.append(candidate_row(year, district, order, names[column], parties[column], total_row[column], percent_row[column]))
    return output


def working_sheet_rows(path: Path) -> list[list[object]]:
    year = 2026
    sheet = openpyxl.load_workbook(path, read_only=True, data_only=True)["Working Data"]
    rows = list(sheet.iter_rows(values_only=True))
    district_header, names, parties = rows[:3]
    columns: dict[int, list[int]] = {district: [] for district in range(1, 53)}
    current_district: int | None = None
    for column in range(2, len(names)):
        if district_header[column] is not None:
            current_district = int(district_header[column])
        if names[column] is not None:
            if current_district is None:
                raise ValueError("2026: candidate column without district")
            columns[current_district].append(column)
    if any(not columns[district] for district in columns):
        raise ValueError("2026: candidate column closure invalid")

    starts = [(index, int(row[0])) for index, row in enumerate(rows[3:], start=3) if row[0] is not None]
    if [district for _, district in starts] != list(range(1, 53)):
        raise ValueError("2026: working-data district row closure invalid")
    output: list[list[object]] = []
    for position, (start, district) in enumerate(starts):
        end = starts[position + 1][0] if position + 1 < len(starts) else len(rows)
        total_row = percent_row = None
        for row in rows[start:end]:
            label = str(row[1]).strip() if row[1] is not None else ""
            if label == "District Totals":
                total_row = row
            elif label == "Percent":
                percent_row = row
        if total_row is None or percent_row is None:
            raise ValueError(f"2026 district {district}: totals missing")
        for order, column in enumerate(columns[district], start=1):
            output.append(candidate_row(year, district, order, names[column], parties[column], total_row[column], percent_row[column]))
    return output


def validate(rows: list[list[object]]) -> None:
    for year, expected in EXPECTED.items():
        cycle = [row for row in rows if row[0] == year]
        if len(cycle) != expected["candidates"] or sum(int(row[8]) for row in cycle) != expected["votes"]:
            raise ValueError(f"{year}: candidate or vote closure invalid")
        if {row[2] for row in cycle} != {f"{district:02d}" for district in range(1, 53)}:
            raise ValueError(f"{year}: district closure invalid")
        for district in range(1, 53):
            contest = [row for row in cycle if row[2] == f"{district:02d}"]
            total = sum(int(row[8]) for row in contest)
            if total < 1 or len({row[4] for row in contest}) != len(contest):
                raise ValueError(f"{year} district {district}: candidate closure invalid")
            for row in contest:
                if abs((int(row[8]) / total) * 1000 - int(row[9])) > 0.500001:
                    raise ValueError(f"{year} district {district}: percentage reconciliation invalid for {row[4]}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    rows = formatted_sheet_rows(args.input_dir / "2022" / "house-statement-of-vote.xlsx", 2022)
    rows += formatted_sheet_rows(args.input_dir / "2024" / "house-statement-of-vote.xlsx", 2024)
    rows += working_sheet_rows(args.input_dir / "2026" / "house-statement-of-vote.xlsx")
    validate(rows)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(HEADER)
        writer.writerows(rows)


if __name__ == "__main__":
    main()
